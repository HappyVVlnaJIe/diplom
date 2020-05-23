import os
import re
import subprocess
from transliterate import translit
import slugify
import docx
import docx2txt
from docx import Document
import openpyxl
doc_directory = ['./students/primaryTemplates',
                 './students/secondaryTemplates',]

doc_settings_path= './server_settings/prompt_and_types.txt'

urls_order_path='./server_settings/urls_order.txt'

primaryTemplates=[]

last_reference="settings"

order=[]

intersection_with_common=dict()

strtype_to_ormtype_dict = {
    'date': 'DateTimeField',
    'string': 'CharField',
    'text': 'TextField',
    'int': 'IntegerField',
    'float': 'FloatField',
}

mandatory_params= {
'date':[] ,
    'string': ["max_length=100"],
    'text':[] ,
    'int': [],
    'float': [],
}

class variable:
    def __init__(self,param):
        self.name=param[0]
        self.type=param[1]
        self.title=param[2]
        self.prompt='Подсказка:'+str.strip(param[3])

def getTextFromDocx(filename):
    doc = docx.Document(filename)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(para.text)
    return '\n'.join(fullText)

def get_template_params():
    template_params = dict()
    template_params_counts = dict()
    for dir in doc_directory:
        for full_filename in os.listdir(dir):
            full_path=os.path.join(dir, full_filename)
            if os.path.isfile(full_path):
                filename, extension = os.path.splitext(full_filename)
                if extension == '.docx' and filename[0]!='~':
                    if (doc_directory.index(dir)==0):
                        primaryTemplates.append(filename)
                    template_params[filename] = set()
                    template_file = docx2txt.process(full_path)
                    template_file = template_file.split('\n')
                    for line in template_file:
                        for param_name in re.findall(r'#([\w_]+)', line):
                            param_name=slugify.slugify(param_name,separator='_')
                            template_params[filename].add( param_name)
                            if param_name not in template_params_counts:
                                template_params_counts[param_name] = 1
                            else:
                                template_params_counts[param_name] += 1

                if extension=='.xlsx'and filename[0]!='~':
                    if (doc_directory.index(dir)==0):
                        primaryTemplates.append(filename)
                    xlsx_file =openpyxl.load_workbook(full_path)
                    sheet=xlsx_file.active
                    template_params[filename] = set()
                    for row in sheet.iter_rows():
                        for cell in row:
                            s=str(cell.value)
                            if (s[0]=='#'):
                                param_name=s.replace('#','')
                                param_name = slugify.slugify(param_name, separator='_')
                                template_params[filename].add(param_name)
                                if param_name not in template_params_counts:
                                    template_params_counts[param_name] = 1
                                else:
                                    template_params_counts[param_name] += 1

    common_template_params = list(filter(lambda param: template_params_counts[param] > 1, template_params_counts))
    for common_template_param in common_template_params:
        for param_name in template_params:
            if param_name not in intersection_with_common:
                intersection_with_common[param_name] = set()
            if common_template_param in template_params[param_name]:
                intersection_with_common[param_name].add(common_template_param)
                template_params[param_name].remove(common_template_param)

    template_params['osnovnye'] = common_template_params
    return template_params

'''
def create_types_and_prompt(template_params_dict):
    with open(settings_path, 'w', encoding='utf8') as template_file:
        for template_name, template_params in template_params_dict.items():
            for param_name in template_params:
                template_file.write('{}|||\n'.format(param_name))
'''

def get_types_and_prompt_dict():
    types_and_prompt=dict()
    with open(doc_settings_path, 'r', encoding='utf8') as template_file:
        for line in template_file:
            param=line.split('|')
            param[0]=slugify.slugify(param[0],separator='_')
            var=variable(param)
            types_and_prompt[var.name]=var

    return types_and_prompt

def generate_models_file(template_params_dict,types_and_prompt_dict):
    with open('users/dynamic_models.py', 'w',encoding='utf8') as models_file:
        models_file.write('from django.db import models\n\n\n')
        for template_name, template_params in template_params_dict.items():
            models_file.write('class {}(models.Model):\n'.format(template_name))
            for param_name in template_params:
                param_type = types_and_prompt_dict[param_name].type
                if not param_type:
                    param_type = 'string'
                if param_type not in strtype_to_ormtype_dict:
                    raise AttributeError("Unknown type: '{}'!".format(param_type))
                mandatory_attrs=[]
                mandatory_attrs.extend(mandatory_params[param_type])
                if template_name=="osnovnye":
                    mandatory_attrs.append("required=False")
                models_file.write('\t{} = models.{}({})\n'.format(param_name, strtype_to_ormtype_dict[param_type],','.join(mandatory_attrs))                                                                  )
            models_file.write('\n\n')

def generate_forms_file(template_params_dict,types_and_prompt_dict):
    with open('users/dynamic_forms.py', 'w',encoding='utf8') as forms_file:
        forms_file.write('from django.forms import ModelForm\n')
        forms_file.write('from django import forms\n')
        forms_file.write('from django.utils.translation import ugettext_lazy as _\n')
        forms_file.write('from .dynamic_models import {}\n\n\n'.format(', '.join(template_params_dict)))
        for template_name, template_params in template_params_dict.items():
            if template_name=='osnovnye':
                continue
            forms_file.write('class {}Form(ModelForm):\n'.format(template_name))
            forms_file.write("\tname='{}'\n".format(translit(template_name,"ru")))
            for field in intersection_with_common[template_name]:
                var=types_and_prompt_dict[field]
                forms_file.write('\t{}=forms.{}(label="{}",help_text="{}")\n'.format(field,strtype_to_ormtype_dict[var.type],var.title,var.prompt))
            forms_file.write('\tclass Meta:\n')
            forms_file.write('\t\tmodel = {}\n'.format(template_name))
            forms_file.write(
                '\t\tfields = ({},)\n\n'.format(', '.join("'{0}'".format(param) for param in template_params)))
            forms_file.write("\t\tlabels = {\n")
            for param in template_params:
                var=types_and_prompt_dict[param]
                forms_file.write("\t\t\t'{}':_('{}'),\n".format(var.name, var.title))
            forms_file.write("\t\t}\n")
            forms_file.write("\t\thelp_texts = {\n")
            for param in template_params:
                var = types_and_prompt_dict[param]
                forms_file.write("\t\t\t'{}':_('{}'),\n".format(var.name, var.prompt))
            forms_file.write("\t\t}\n")

        """for template_name, template_params in intersection_with_common.items():
            forms_file.write('class {}Form(ModelForm):\n'.format(template_name))
            forms_file.write("\tname='{}'\n".format("Intersection"+translit(template_name, "ru")))
            forms_file.write('\tclass Meta:\n')
            forms_file.write('\t\tmodel = {}\n'.format(template_name))
            forms_file.write(
                '\t\tfields = ({},)\n\n'.format(', '.join("'{0}'".format(param) for param in template_params)))
            forms_file.write("\t\tlabels = {\n")
            for param in template_params:
                var = types_and_prompt_dict[param]
                forms_file.write("\t\t\t'{}':_('{}'),\n".format(var.name, var.title))
            forms_file.write("\t\t}\n")
            forms_file.write("\t\thelp_texts = {\n")
            for param in template_params:
                var = types_and_prompt_dict[param]
                forms_file.write("\t\t\t'{}':_('{}'),\n".format(var.name, var.prompt))
            forms_file.write("\t\t}\n")"""

def generate_admin_file(template_params_dict):
    with open('users/admin.py', 'w') as admin_file:
        admin_file.write('from django.contrib import admin\n')
        admin_file.write('from .dynamic_models import {}\n\n\n'.format(', '.join(template_params_dict)))
        for template_name in template_params_dict:
            admin_file.write('admin.site.register({})\n'.format(template_name))

def generate_views_file(template_params_dict):
    with open('users/dynamic_views.py', 'w',encoding='utf8') as views_file:
        views_file.write('from .dynamic_forms import *\n')
        views_file.write('from django.contrib.auth.decorators import login_required\n'
                         'from django.shortcuts import render, redirect\n\n')
        for i in range(1,len(order)):
            views_file.write('@login_required\n')
            views_file.write('def stage{}(request):\n'.format(i+1))
            views_file.write('\tunique_forms=[]\n')
            views_file.write("\tif request.method == 'POST':\n")
            for form_name in order[i].split('|'):
                form_name=str.strip(form_name)
                views_file.write("\t\tunique_forms.append({}Form(data=request.POST))\n".format(form_name))
                views_file.write("\t\tif unique_forms[-1].is_valid():\n")
                views_file.write("\t\t\tunique_forms[-1].save()\n")
            if i==len(order)-1:
                views_file.write("\t\treturn redirect('/{}')\n".format(last_reference))
            else:
                views_file.write("\t\treturn redirect('/stage{}')\n".format(i + 2))
            views_file.write("\telse:\n")
            for form_name in order[i].split('|'):
                form_name = str.strip(form_name)
                views_file.write("\t\tunique_forms.append({}Form())\n".format(form_name))
                views_file.write('\t\treturn render(request,"profile/stage%d.html",{"forms": unique_forms})\n\n'%(i+1))

def generate_order():
    with open(urls_order_path, 'r') as order_file:
        for line in order_file:
            order.append(line)

def generate_urls_file():
    with open("users/dynamic_urls.py", 'w') as dynamic_urls_file:
        dynamic_urls_file.write("from django.conf.urls import url\n")
        dynamic_urls_file.write("from django.urls import path\n")
        dynamic_urls_file.write("from .dynamic_views import *\n")
        dynamic_urls_file.write("urlpatterns = [\n")
        for i in range(len(order)):
            dynamic_urls_file.write("\tpath('stage{0}', stage{0}, name='stage{0}_view'),\n".format(i+1))
        dynamic_urls_file.write("]\n")

def generate_html_files():
    for i in range(len(order)):
        with open("templates/profile/stage{}.html".format(i+1), 'w',encoding='utf8') as html_file:
            html_file.write('{% extends "base_generic.html" %}\n')
            html_file.write('{% block content %}\n')
            html_file.write('{% for form in forms %}')
            html_file.write('<form action="" method="post">{% csrf_token %}\n')
            html_file.write('{{ form.as_p }}\n')
            html_file.write('<input type="submit" value="Сохранить">\n')
            html_file.write('</form>\n')
            html_file.write('{% endfor %}')
            html_file.write('{% endblock %}\n')


fetched_template_params = get_template_params()

#create_types_and_prompt(fetched_template_params)

types_and_prompt_dict=get_types_and_prompt_dict()
generate_order()
generate_html_files()
generate_models_file(fetched_template_params,types_and_prompt_dict)
generate_forms_file(fetched_template_params,types_and_prompt_dict)
generate_admin_file(fetched_template_params)
generate_views_file(fetched_template_params)
generate_urls_file()
make_migrations_retcode = subprocess.run(["python", "manage.py", "makemigrations","--noinput"], capture_output=True).returncode
migrate_retcode = -1
if make_migrations_retcode == 0:
    migrate_retcode = subprocess.run(["python", "manage.py", "migrate"], capture_output=True).returncode
"""
if migrate_retcode == 0:
    result = subprocess.run(["python", "manage.py", "runserver"], capture_output=True)
    print(result)
"""